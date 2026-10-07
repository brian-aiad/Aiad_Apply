from pathlib import Path

import pytest
from aiadapply_v2 import pipeline
from aiadapply_v2.documents.model import parse_resume_docx
from aiadapply_v2.documents.writer import write_resume_candidate
from aiadapply_v2.evidence.loaders import build_evidence_graph
from aiadapply_v2.grading.keywords import NON_PROSE_QUALIFICATIONS, grade_job_keywords
from aiadapply_v2.parsers.linkedin_simplify import parse_linkedin_simplify
from aiadapply_v2.pipeline import (
    _certification_only_terms,
    _display_skill,
    _enforce_export_boundaries,
    _ensure_important_keyword_placement,
    _ensure_target_identity,
    _ensure_transferable_review_risks,
    _fuse_skill_inventory,
    _measured_skill_character_budget,
    _merge_transferability_map,
    _normalize_acronym_redundancy,
    _normalize_claim_risk_placements,
    _normalize_prose_collocations,
    _order_rewrite_variants,
    _place_direct_category_keyword,
    _prune_absent_claim_risks,
    _prune_direct_evidence_risks,
    _prune_resolved_nonexportable_risks,
    _remove_unproven_reproduction,
    _repair_equal_length_fallbacks,
    _repair_skill_row_overflow,
    _restore_unjustified_shortening,
    _sanitize_shorter_fallbacks,
    _select_compression_candidate,
    _select_expansion_candidate,
    _select_overflow_reversion_candidate,
    _supported_keywords,
    _use_rendered_skill_rows,
    transform_resume,
)
from aiadapply_v2.planning.rewrite_plan import all_proposed_text
from aiadapply_v2.profiling.role_profile import build_target_role_profile
from aiadapply_v2.schemas import (
    ClaimRisk,
    EvidenceMatch,
    EvidenceStrength,
    JobKeyword,
    KeywordKind,
    KeywordPriority,
    LayoutResult,
    ReasoningResult,
    RiskLevel,
    TargetRoleProfile,
    TransferabilityMap,
)
from aiadapply_v2.semantic.matcher import LexicalSemanticEncoder, build_transferability_map
from aiadapply_v2.text import split_skill_values
from aiadapply_v2.validation.resume import validate_rewrite_plan

from .helpers import IdentityReasoner, identity_plan

BASE = Path("data/resumes/Brian_Aiad_BASE.docx")


def test_project_api_wording_survives_plural_only_rewrite() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = next(
        item for item in plan.bullets if item.paragraph_id == "projects.loavenly.bullet.2"
    )
    source = bullet.text
    assert "API," in source
    bullet.text = source.replace("API,", "APIs,")
    bullet.shorter_text = bullet.text

    _restore_unjustified_shortening(plan, document, [])

    assert bullet.text == source


@pytest.mark.parametrize("corrupt_optimized_text", [False, True])
def test_end_to_end_pipeline_exports_job_specific_names_after_validation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, corrupt_optimized_text: bool
) -> None:
    raw = """
Home
Jobs
Company logo for, Example Systems.
Example Systems
Technical Support Engineer
Los Angeles, CA · Posted today
Hybrid
Full-time
About the job
Example Systems provides production software to business customers. This paragraph
adds realistic page noise and sufficient text for a complete paste parser fixture.
Responsibilities
Troubleshoot production incidents using SQL, Postman, Jira, log analysis, and API debugging.
Resolve technical support tickets and document root cause analysis for SaaS applications.
Required Qualifications
Experience with SQL, Postman, Jira, technical support, troubleshooting, and REST APIs.
Strong written communication and cross-functional support experience.
Preferred Qualifications
Experience with Microsoft 365, Python, PowerShell, Linux, and Confluence.
Set alert for similar jobs
Candidates who clicked apply
100 total
About
Accessibility
LinkedIn Corporation
""" + ("navigation noise " * 20)
    progress: list[str] = []
    # This identity-plan fixture tests render/export integrity, not model writing.
    # Broad rewrite enforcement has its own unit cases and real-listing verification.
    monkeypatch.setattr(pipeline, "tailoring_feedback", lambda *args, **kwargs: [])
    if corrupt_optimized_text:
        from docx import Document

        original_optimizer = pipeline.write_ats_optimized_docx

        def corrupt_optimizer(source: Path, destination: Path) -> Path:
            result = original_optimizer(source, destination)
            document = Document(result)
            for paragraph in document.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace("SQL", "spreadsheets")
            document.save(result)
            return result

        monkeypatch.setattr(pipeline, "write_ats_optimized_docx", corrupt_optimizer)
        with pytest.raises(pipeline.TransformationError, match="changed the validated resume text"):
            transform_resume(
                raw_paste=raw,
                base_resume=BASE,
                output_dir=tmp_path,
                reasoner=IdentityReasoner(),
                semantic_encoder=LexicalSemanticEncoder(),
                aggressive_draft=False,
            )
        assert not (tmp_path / "transformation_report.json").exists()
        return
    report = transform_resume(
        raw_paste=raw,
        base_resume=BASE,
        output_dir=tmp_path,
        reasoner=IdentityReasoner(),
        semantic_encoder=LexicalSemanticEncoder(),
        progress=progress.append,
        aggressive_draft=False,
    )

    assert (
        report.output_docx.name
        == "Brian_Aiad_Resume_Example_Systems_Technical_Support_Engineer.docx"
    )
    assert (
        report.output_pdf.name == "Brian_Aiad_Resume_Example_Systems_Technical_Support_Engineer.pdf"
    )
    assert report.output_docx.exists()
    assert report.output_pdf.exists()
    assert report.validation.passed
    assert report.layout.passed
    assert report.layout.page_count == 1
    assert report.keyword_decisions
    assert report.changes
    assert all(change.before_text for change in report.changes)
    assert {change.paragraph_id for change in report.changes} == set(
        parse_resume_docx(BASE).editable_paragraph_ids
    )
    assert (tmp_path / "transformation_report.json").exists()
    assert (tmp_path / "transformation_report.md").exists()
    audit = tmp_path / "character_audit.json"
    assert audit.exists()
    assert '"format_integrity_passed": true' in audit.read_text(encoding="utf-8")
    assert progress[0] == "Parsing and grading the job posting"
    assert progress[-1] == "Transformation complete"


def test_target_identity_is_repaired_in_primary_and_fallback_summaries() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/recent_floqast_integrations_support_2026_08_05.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(),
        rewrite_plan=identity_plan(document),
    )

    _ensure_target_identity(reasoning)

    assert reasoning.rewrite_plan.professional_identity == "Technical Support Engineer"
    assert reasoning.role_profile.professional_identity == "Technical Support Engineer"
    assert reasoning.rewrite_plan.summary.text.startswith("Technical Support Engineer with")
    assert reasoning.rewrite_plan.summary.shorter_text.startswith("Technical Support Engineer with")


def test_erp_identity_does_not_claim_an_unsupported_named_system() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/liquid_iv_d365_technical_analyst_2026_08_12.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(
            unsupported_terms=["D365", "D365 F&O", "Dynamics 365", "ERP"]
        ),
        rewrite_plan=identity_plan(parse_resume_docx(BASE)),
    )

    _ensure_target_identity(reasoning)

    assert reasoning.rewrite_plan.professional_identity == "Application Support Analyst"
    assert reasoning.rewrite_plan.summary.text.startswith("Application Support Analyst with")
    assert not any(
        term in reasoning.rewrite_plan.summary.text for term in ("D365", "Dynamics 365", "ERP")
    )


def test_deterministic_direct_evidence_fills_incomplete_reasoner_map() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/recent_cresta_application_support_2026_08_05.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    profile = build_target_role_profile(job, keywords)
    graph = build_evidence_graph(document, keywords)
    preliminary = build_transferability_map(
        profile,
        keywords,
        graph,
        LexicalSemanticEncoder(),
    )
    reasoning = ReasoningResult(
        role_profile=profile,
        transferability_map=TransferabilityMap(),
        rewrite_plan=identity_plan(document),
    )

    _merge_transferability_map(reasoning, preliminary, keywords)

    by_term = {
        match.target_term.casefold(): match for match in reasoning.transferability_map.matches
    }
    assert by_term["apis"].strength == "direct"
    assert by_term["api integrations"].strength == "direct"
    assert by_term["h.323"].strength == "unsupported"
    assert by_term["zendesk"].strength == "unsupported"
    assert "APIs" in reasoning.transferability_map.direct_terms
    assert len(reasoning.transferability_map.matches) == len(
        [keyword for keyword in keywords if keyword.accepted]
    )


def test_model_cannot_promote_formal_foundry_requirements_without_source_evidence() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/mccarthy_product_analyst_ii_foundry_2026_09_21.txt").read_text()
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    profile = build_target_role_profile(job, keywords)
    preliminary = build_transferability_map(
        profile, keywords, build_evidence_graph(document, keywords), LexicalSemanticEncoder()
    )
    protected = {"runbooks", "acceptance criteria", "user stories", "tier 2 support", "foundry"}
    promoted = [
        match.model_copy(update={"strength": EvidenceStrength.direct})
        for match in preliminary.matches
        if match.target_term.casefold() in protected
    ]
    assert len(promoted) == len(protected)
    reasoning = ReasoningResult(
        role_profile=profile,
        transferability_map=TransferabilityMap(matches=promoted),
        rewrite_plan=identity_plan(document),
    )
    _merge_transferability_map(reasoning, preliminary, keywords)
    for match in reasoning.transferability_map.matches:
        if match.target_term.casefold() in protected:
            assert match.strength in {
                EvidenceStrength.weakly_transferable,
                EvidenceStrength.unsupported,
            }


def test_curated_transfer_bridge_is_stable_when_reasoner_downgrades_it() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/liquid_iv_d365_technical_analyst_2026_08_12.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    profile = build_target_role_profile(job, keywords)
    preliminary = build_transferability_map(
        profile,
        keywords,
        build_evidence_graph(document, keywords),
        LexicalSemanticEncoder(),
    )
    deterministic = next(
        match
        for match in preliminary.matches
        if match.target_term.casefold() == "cross-functional collaboration"
    )
    assert deterministic.strength == EvidenceStrength.strongly_transferable
    downgraded = deterministic.model_copy(update={"strength": EvidenceStrength.unsupported})
    reasoning = ReasoningResult(
        role_profile=profile,
        transferability_map=TransferabilityMap(
            matches=[downgraded],
            unsupported_terms=[downgraded.target_term],
        ),
        rewrite_plan=identity_plan(document),
    )

    _merge_transferability_map(reasoning, preliminary, keywords)

    merged = next(
        match
        for match in reasoning.transferability_map.matches
        if match.target_term.casefold() == "cross-functional collaboration"
    )
    assert merged.strength == EvidenceStrength.strongly_transferable


def test_protected_certifications_and_suite_aliases_count_as_existing_evidence() -> None:
    document = parse_resume_docx(BASE)
    encoder = LexicalSemanticEncoder()

    la_job = parse_linkedin_simplify(
        Path("data/fixtures/los_angeles_times_application_support_full.txt").read_text(
            encoding="utf-8"
        )
    )
    la_keywords = grade_job_keywords(la_job)
    la_profile = build_target_role_profile(la_job, la_keywords)
    la_graph = build_evidence_graph(document, la_keywords)
    la_map = build_transferability_map(la_profile, la_keywords, la_graph, encoder)
    azure = next(match for match in la_map.matches if match.target_term.casefold() == "azure")

    assert azure.strength == EvidenceStrength.direct
    assert azure.evidence_id == "evidence.certifications.protected_body.40"

    niagara_job = parse_linkedin_simplify(
        Path("data/fixtures/niagara_it_systems_functional_i_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    niagara_keywords = grade_job_keywords(niagara_job)
    niagara_profile = build_target_role_profile(niagara_job, niagara_keywords)
    niagara_graph = build_evidence_graph(document, niagara_keywords)
    niagara_map = build_transferability_map(
        niagara_profile,
        niagara_keywords,
        niagara_graph,
        encoder,
    )
    office = next(
        match for match in niagara_map.matches if match.target_term.casefold() == "microsoft office"
    )

    assert office.strength == EvidenceStrength.direct


def test_existing_protected_keyword_is_not_duplicated_into_skills() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/los_angeles_times_application_support_full.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    azure = next(keyword for keyword in keywords if keyword.normalized == "azure")
    document = parse_resume_docx(BASE)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(direct_terms=["Azure"]),
        rewrite_plan=identity_plan(document),
    )

    _ensure_important_keyword_placement(reasoning, [azure], document)

    skill_text = " ".join(
        skill for line in reasoning.rewrite_plan.skills.lines for skill in line.skills
    )
    assert "Azure" not in skill_text


def test_stem_degree_is_satisfied_by_education_not_inserted_as_a_skill() -> None:
    job = parse_linkedin_simplify(
        Path(
            "data/fixtures/rtx_raytheon_systems_engineer_radar_integration_test_2026_08_10.txt"
        ).read_text(encoding="utf-8")
    )
    keywords = grade_job_keywords(job)
    stem = next(keyword for keyword in keywords if keyword.normalized == "stem degree")
    document = parse_resume_docx(BASE)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(direct_terms=["STEM degree"]),
        rewrite_plan=identity_plan(document),
    )
    technical = next(
        line
        for line in reasoning.rewrite_plan.skills.lines
        if line.paragraph_id == "skills.technical_support"
    )
    technical.skills.append("STEM degree")

    _fuse_skill_inventory(
        reasoning.rewrite_plan,
        document,
        keywords,
        skill_ineligible_terms=NON_PROSE_QUALIFICATIONS,
    )
    _ensure_important_keyword_placement(reasoning, [stem], document)

    skill_text = " ".join(
        skill for line in reasoning.rewrite_plan.skills.lines for skill in line.skills
    )
    assert "STEM degree" not in skill_text


def test_credential_specializations_require_exact_evidence_while_stem_degree_is_direct() -> None:
    document = parse_resume_docx(BASE)

    radar_job = parse_linkedin_simplify(
        Path(
            "data/fixtures/rtx_raytheon_systems_engineer_radar_integration_test_2026_08_10.txt"
        ).read_text(encoding="utf-8")
    )
    radar_keywords = grade_job_keywords(radar_job)
    radar_map = build_transferability_map(
        build_target_role_profile(radar_job, radar_keywords),
        radar_keywords,
        build_evidence_graph(document, radar_keywords),
        LexicalSemanticEncoder(),
    )
    radar_matches = {match.target_term.casefold(): match for match in radar_map.matches}
    assert radar_matches["stem degree"].strength == EvidenceStrength.direct
    assert radar_matches["security clearance"].strength == EvidenceStrength.unsupported

    rf_job = parse_linkedin_simplify(
        Path("data/fixtures/rtx_raytheon_rf_microwave_antenna_engineer_i_01864246.txt").read_text(
            encoding="utf-8"
        )
    )
    rf_keywords = grade_job_keywords(rf_job)
    rf_map = build_transferability_map(
        build_target_role_profile(rf_job, rf_keywords),
        rf_keywords,
        build_evidence_graph(document, rf_keywords),
        LexicalSemanticEncoder(),
    )
    rf_matches = {match.target_term.casefold(): match for match in rf_map.matches}
    assert rf_matches["electrical engineering"].strength == EvidenceStrength.unsupported
    assert rf_matches["u.s. citizenship"].strength == EvidenceStrength.unsupported


def test_supported_stretch_role_actions_are_placed_in_existing_investigation_evidence() -> None:
    job = parse_linkedin_simplify(
        Path(
            "data/fixtures/rtx_raytheon_semiconductor_manufacturing_engineer_01862457.txt"
        ).read_text(encoding="utf-8")
    )
    keywords = grade_job_keywords(job)
    supported = [
        keyword
        for keyword in keywords
        if keyword.normalized in {"problem-solving", "root cause and corrective action"}
    ]
    document = parse_resume_docx(BASE)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(
            strongly_transferable_terms=[keyword.term for keyword in supported]
        ),
        rewrite_plan=identity_plan(document),
    )

    _ensure_important_keyword_placement(reasoning, supported, document)

    investigation = next(
        bullet
        for bullet in reasoning.rewrite_plan.bullets
        if bullet.paragraph_id == "experience.original_insurance.bullet.2"
    )
    assert investigation.text.startswith("Investigated production incidents")
    assert "Applied problem-solving" not in investigation.text
    assert "root cause and corrective action" in investigation.text


def test_certification_only_system_is_removed_from_proposed_skill_additions() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/los_angeles_times_application_support_full.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    profile = build_target_role_profile(job, keywords)
    graph = build_evidence_graph(document, keywords)
    transferability = build_transferability_map(
        profile,
        keywords,
        graph,
        LexicalSemanticEncoder(),
    )
    plan = identity_plan(document)
    cloud = next(line for line in plan.skills.lines if line.paragraph_id == "skills.cloud_systems")
    cloud.skills.append("Azure")

    _fuse_skill_inventory(
        plan,
        document,
        keywords,
        skill_ineligible_terms=_certification_only_terms(transferability),
    )

    assert "Azure" not in cloud.skills


def test_sla_acronym_redundancy_is_rewritten_as_natural_prose() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    plan.summary.text = "Resolved tickets to support SLA and Service Level Agreements."
    plan.summary.shorter_text = plan.summary.text

    _normalize_acronym_redundancy(plan)

    assert plan.summary.text == "Resolved tickets to meet Service Level Agreements (SLA) targets."
    assert plan.summary.shorter_text == plan.summary.text


def test_direct_web_services_category_is_placed_through_existing_webhook_evidence() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/los_angeles_times_application_support_full.txt").read_text(
            encoding="utf-8"
        )
    )
    keyword = next(item for item in grade_job_keywords(job) if item.normalized == "web services")
    plan = identity_plan(parse_resume_docx(BASE))

    _place_direct_category_keyword(plan, keyword)

    bullet = next(
        item
        for item in plan.bullets
        if item.paragraph_id == "experience.original_insurance.bullet.5"
    )
    assert "webhook-based web services" in bullet.text
    assert "webhook-based web services" in bullet.shorter_text


def test_monitoring_does_not_mechanically_rewrite_an_investigation() -> None:
    monitoring = JobKeyword(
        term="monitoring",
        normalized="monitoring",
        kind=KeywordKind.action,
        priority=KeywordPriority.high,
        hiring_importance=70,
        placement_utility=75,
    )
    plan = identity_plan(parse_resume_docx(BASE))

    _place_direct_category_keyword(plan, monitoring)

    bullet = next(
        item
        for item in plan.bullets
        if item.paragraph_id == "experience.original_insurance.bullet.2"
    )
    assert bullet.text.startswith("Investigated production incidents")
    assert "using log analysis, SQL queries, and API debugging" in bullet.text


def test_business_analysis_is_not_mechanically_prefixed_to_a_project() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/niagara_it_systems_functional_i_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    keyword = next(
        item for item in grade_job_keywords(job) if item.normalized == "business analysis"
    )
    plan = identity_plan(parse_resume_docx(BASE))

    _place_direct_category_keyword(plan, keyword)

    bullet = next(
        item for item in plan.bullets if item.paragraph_id == "experience.wehelp.bullet.1"
    )
    assert bullet.text.startswith("Identified the need")
    assert "independently built, tested, and deployed Loavenly" in bullet.text
    assert "staff and volunteers" in bullet.text


def test_client_operations_terms_are_placed_in_supported_volunteer_evidence() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/nesco_technical_client_operations_specialist_2026_08_12.txt").read_text(
            encoding="utf-8"
        )
    )
    by_term = {item.normalized: item for item in grade_job_keywords(job)}
    plan = identity_plan(parse_resume_docx(BASE))

    _place_direct_category_keyword(plan, by_term["software implementation"])
    _place_direct_category_keyword(plan, by_term["software adoption"])

    implementation = next(
        item for item in plan.bullets if item.paragraph_id == "experience.wehelp.bullet.1"
    )
    adoption = next(
        item for item in plan.bullets if item.paragraph_id == "experience.wehelp.bullet.2"
    )
    assert implementation.text.startswith("Led software implementation")
    assert "Loavenly by independently building, testing, and deploying" in implementation.text
    assert "software adoption" in adoption.text
    assert "software adoption" in adoption.shorter_text


def test_support_process_terms_are_placed_in_direct_experience_evidence() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/liquid_iv_d365_technical_analyst_2026_08_12.txt").read_text(
            encoding="utf-8"
        )
    )
    by_term = {item.normalized: item for item in grade_job_keywords(job)}
    plan = identity_plan(parse_resume_docx(BASE))

    _place_direct_category_keyword(plan, by_term["ticketing"])
    _place_direct_category_keyword(plan, by_term["incident management"])

    ticket = next(
        item
        for item in plan.bullets
        if item.paragraph_id == "experience.original_insurance.bullet.1"
    )
    incident = next(
        item for item in plan.bullets if item.paragraph_id == "experience.csulb.bullet.3"
    )
    assert "Jira ticketing" in ticket.text
    assert "Jira ticketing" in ticket.shorter_text
    assert incident.text.startswith("Supported incident management")
    assert "campus IT and third-party vendors" in incident.text


def test_account_management_is_placed_in_existing_identity_administration_evidence() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/liquid_iv_d365_technical_analyst_2026_08_12.txt").read_text(
            encoding="utf-8"
        )
    )
    keyword = next(
        item for item in grade_job_keywords(job) if item.normalized == "account management"
    )
    plan = identity_plan(parse_resume_docx(BASE))
    source_bullet = next(
        item
        for item in plan.bullets
        if item.paragraph_id == "experience.original_insurance.bullet.4"
    )
    source_length = len(source_bullet.text)

    _place_direct_category_keyword(plan, keyword)

    bullet = next(
        item
        for item in plan.bullets
        if item.paragraph_id == "experience.original_insurance.bullet.4"
    )
    assert "MFA, RBAC, and account management" in bullet.text
    assert "MFA, RBAC, and account management" in bullet.shorter_text
    assert len(bullet.text) == source_length
    assert "account management" in bullet.target_terms


def test_weakly_transferable_terms_are_not_export_supported() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/los_angeles_times_application_support_full.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(weakly_transferable_terms=["change management"]),
        rewrite_plan=identity_plan(parse_resume_docx(BASE)),
    )

    supported = _supported_keywords(reasoning, keywords)

    assert "change management" not in {keyword.normalized for keyword in supported}


def test_literal_keyword_collocations_are_rewritten_as_natural_prose() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = next(item for item in plan.bullets if item.paragraph_id == "experience.csulb.bullet.3")
    bullet.text = (
        "Maintained documentation while coordinating cross-functional collaboration with campus IT."
    )
    volunteer = next(
        item for item in plan.bullets if item.paragraph_id == "experience.wehelp.bullet.2"
    )
    volunteer.text = (
        "Support live distributions through user access management, troubleshooting issues, "
        "update validation, customer service, and training for staff."
    )

    _normalize_prose_collocations(plan)

    assert "supporting cross-functional collaboration" in bullet.text
    assert "by managing user access" in volunteer.text
    assert "delivering customer service and training to staff" in volunteer.text

    bullet.text = (
        "Maintained incident records, demonstrating organizational skills while coordinating "
        "with IT."
    )
    _normalize_prose_collocations(plan)
    assert bullet.text == (
        "Applied organizational skills to maintain incident records while coordinating with IT."
    )

    platform = next(
        item for item in plan.bullets if item.paragraph_id == "projects.loavenly.bullet.1"
    )
    platform.text = platform.text.replace(
        "across three food bank locations",
        ", and system maintenance across three food bank locations",
    )
    _normalize_prose_collocations(plan)
    assert "system maintenance" not in platform.text.casefold()

    plan.summary.text = "Resolved tickets by applying diagnostics, and Service Level Agreements."
    _normalize_prose_collocations(plan)
    assert plan.summary.text == (
        "Resolved tickets by applying diagnostics while meeting Service Level Agreements."
    )

    plan.summary.text = (
        "Resolved tickets by applying diagnostics and root cause analysis within "
        "Service Level Agreements."
    )
    _normalize_prose_collocations(plan)
    assert plan.summary.text == (
        "Resolved tickets by applying diagnostics and root cause analysis while meeting "
        "Service Level Agreements."
    )

    volunteer.text = (
        "Support Loavenly during live distributions through customer service, "
        "user-access management, troubleshooting production issues, validating updates, "
        "and training staff and volunteers."
    )
    _normalize_prose_collocations(plan)
    assert volunteer.text == (
        "Provide customer service for Loavenly during live distributions by managing user "
        "access, troubleshooting production issues, validating updates, and training staff "
        "and volunteers."
    )

    bullet.text = (
        "Supported incident management through documentation, escalation notes, and recurring "
        "issue tracking, using cross-functional collaboration and written and verbal "
        "communication skills with campus IT and third-party vendors."
    )
    _normalize_prose_collocations(plan)
    assert bullet.text == (
        "Supported incident management by maintaining documentation, escalation notes, and "
        "recurring issue tracking while applying written and verbal communication skills during "
        "cross-functional collaboration with campus IT and third-party vendors."
    )

    ticket = next(
        item
        for item in plan.bullets
        if item.paragraph_id == "experience.original_insurance.bullet.1"
    )
    ticket.text = "Resolved tickets across Microsoft 365,, maintaining SLA performance."
    _normalize_prose_collocations(plan)
    assert ",," not in ticket.text

    volunteer.text = (
        "Support Loavenly during live distributions through customer service, "
        "user-access management, production troubleshooting, update validation, and "
        "training for staff and volunteers."
    )
    _normalize_prose_collocations(plan)
    assert volunteer.text == (
        "Provide customer service for Loavenly during live distributions by managing user "
        "access, troubleshooting production issues, validating updates, and training staff "
        "and volunteers."
    )

    volunteer.text = (
        "Support Loavenly during live distributions through customer service, "
        "user-access management, troubleshooting production issues, update validation, "
        "and training for staff and volunteers."
    )
    _normalize_prose_collocations(plan)
    assert volunteer.text == (
        "Provide customer service for Loavenly during live distributions by managing user "
        "access, troubleshooting production issues, validating updates, and training staff "
        "and volunteers."
    )


def test_soft_skills_do_not_replace_concrete_action_verbs() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/niagara_it_systems_functional_i_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    soft_terms = {"critical thinking", "problem-solving", "task prioritization"}
    keywords = [keyword for keyword in grade_job_keywords(job) if keyword.normalized in soft_terms]
    document = parse_resume_docx(BASE)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(),
        rewrite_plan=identity_plan(document),
    )

    _ensure_important_keyword_placement(reasoning, keywords, document)

    proposed_text = "\n".join(all_proposed_text(reasoning.rewrite_plan)).casefold()
    for term in soft_terms:
        assert term not in proposed_text
    incident = next(
        item
        for item in reasoning.rewrite_plan.bullets
        if item.paragraph_id == "experience.original_insurance.bullet.2"
    )
    ticket = next(
        item
        for item in reasoning.rewrite_plan.bullets
        if item.paragraph_id == "experience.original_insurance.bullet.1"
    )
    assert incident.text.startswith("Investigated production incidents")
    assert ticket.text.startswith("Resolved 200+")


def test_communication_keywords_do_not_force_filler_into_evidence() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/liquid_iv_d365_technical_analyst_2026_08_12.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(),
        rewrite_plan=identity_plan(document),
    )
    incident = next(item for item in keywords if item.normalized == "incident management")
    _place_direct_category_keyword(reasoning.rewrite_plan, incident)

    supported = [
        item
        for item in keywords
        if item.normalized in {"cross-functional collaboration", "verbal communication"}
    ]
    _ensure_important_keyword_placement(reasoning, supported, document)

    bullet = next(
        item
        for item in reasoning.rewrite_plan.bullets
        if item.paragraph_id == "experience.csulb.bullet.3"
    )
    for text in (bullet.text, bullet.shorter_text):
        assert "coordinating with campus IT and third-party vendors" in text
        assert "while using verbal communication" not in text


def test_encompass_role_separates_platform_gaps_from_transferable_support_work() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/bank_of_hope_encompass_full_2026_08_06.txt").read_text(encoding="utf-8")
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    profile = build_target_role_profile(job, keywords)
    graph = build_evidence_graph(document, keywords)
    transferability = build_transferability_map(
        profile,
        keywords,
        graph,
        LexicalSemanticEncoder(),
    )
    by_term = {match.target_term.casefold(): match for match in transferability.matches}

    assert by_term["encompass"].strength == "unsupported"
    assert by_term["loan origination system"].strength == "unsupported"
    assert by_term["residential lending"].strength == "unsupported"
    assert by_term["troubleshooting"].strength == "direct"
    assert by_term["system testing"].strength == "strongly_transferable"
    assert by_term["vendor interfaces"].strength == "strongly_transferable"
    assert by_term["workflow documentation"].strength == "strongly_transferable"


def test_rewrite_plan_rejects_unsupported_protocol_in_resume_prose() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/recent_cresta_application_support_2026_08_05.txt").read_text(
            encoding="utf-8"
        )
    )
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    plan.summary.text += " with H.323 support"

    validation = validate_rewrite_plan(
        document,
        plan,
        [],
        build_target_role_profile(job, grade_job_keywords(job)),
        forbidden_terms=["H.323"],
    )

    assert not validation.passed
    assert any(issue.code == "unsupported_keyword_inserted" for issue in validation.issues)


def test_gap_risk_is_normalized_when_term_is_present_in_resume() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/recent_cresta_application_support_2026_08_05.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    technical_support = next(
        line for line in plan.skills.lines if line.paragraph_id == "skills.technical_support"
    )
    technical_support.skills.append("Web Hosting")
    plan.claim_risks.append(
        ClaimRisk(
            claim="Web-hosting experience",
            target_requirement="Experience with web hosting.",
            strength=EvidenceStrength.weakly_transferable,
            risk_level=RiskLevel.medium,
            explanation="Excluded from resume prose.",
            selected_placement="Recorded as a gap; omitted from resume prose.",
        )
    )
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(),
        rewrite_plan=plan,
    )

    _normalize_claim_risk_placements(reasoning, keywords)

    risk = reasoning.rewrite_plan.claim_risks[0]
    assert risk.selected_placement == "skills.technical_support"
    assert "verify this wording" in risk.explanation


def test_free_form_risk_placement_is_normalized_before_export_enforcement() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/pds_health_epic_analyst_full_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    customer_service = next(
        keyword for keyword in keywords if keyword.normalized == "customer service"
    )
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    technical = next(
        line for line in plan.skills.lines if line.paragraph_id == "skills.technical_support"
    )
    technical.skills.append("Customer Service")
    risk = ClaimRisk(
        claim="Customer Service proficiency",
        target_requirement="customer service",
        strength=EvidenceStrength.strongly_transferable,
        risk_level=RiskLevel.medium,
        explanation="Review before export.",
        selected_placement="Use this in the support skills row.",
        export_allowed=False,
    )
    plan.claim_risks.append(risk)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(),
        rewrite_plan=plan,
    )

    _normalize_claim_risk_placements(reasoning, [customer_service])
    _enforce_export_boundaries(plan, document)

    assert risk.selected_placement == "skills.technical_support"
    assert "Customer Service" not in technical.skills


def test_model_risk_is_removed_when_deterministic_evidence_is_direct() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    plan.claim_risks.append(
        ClaimRisk(
            claim="Microsoft Office proficiency",
            target_requirement="Microsoft Office",
            strength=EvidenceStrength.strongly_transferable,
            risk_level=RiskLevel.medium,
            explanation="Model was overly cautious.",
            selected_placement="skills.tools",
            export_allowed=False,
        )
    )
    reasoning = ReasoningResult(
        role_profile=TargetRoleProfile(
            title="IT Systems Functional I",
            normalized_role_family="business_systems_functional",
            professional_identity="Business Systems Support Analyst",
        ),
        transferability_map=TransferabilityMap(
            matches=[
                EvidenceMatch(
                    target_term="Microsoft Office",
                    target_requirement="Microsoft Office proficiency",
                    evidence_id="evidence.skills.tools",
                    source_text="Microsoft 365 and Excel",
                    semantic_score=1.0,
                    action_compatibility=0.5,
                    system_compatibility=0.9,
                    environment_compatibility=0.1,
                    outcome_compatibility=0.1,
                    strength=EvidenceStrength.direct,
                    reasoning="Direct suite evidence.",
                    suggested_placement="skills.tools",
                )
            ],
            direct_terms=["Microsoft Office"],
        ),
        rewrite_plan=plan,
    )

    _prune_direct_evidence_risks(reasoning)

    assert reasoning.rewrite_plan.claim_risks == []


def test_sensitive_transferable_duty_gets_an_automatic_review_risk() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/niagara_it_systems_functional_i_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    profile = build_target_role_profile(job, keywords)
    graph = build_evidence_graph(document, keywords)
    transferability = build_transferability_map(
        profile,
        keywords,
        graph,
        LexicalSemanticEncoder(),
    )
    plan = identity_plan(document)
    bullet = next(
        item for item in plan.bullets if item.paragraph_id == "experience.wehelp.bullet.1"
    )
    bullet.text = bullet.text.replace(
        "Identified the need",
        "Applied business analysis and identified the need",
    )
    reasoning = ReasoningResult(
        role_profile=profile,
        transferability_map=transferability,
        rewrite_plan=plan,
    )

    _ensure_transferable_review_risks(reasoning, keywords, document)

    risk = next(risk for risk in plan.claim_risks if risk.claim == "business analysis")
    assert risk.strength == EvidenceStrength.strongly_transferable
    assert risk.risk_level == RiskLevel.medium
    assert risk.selected_placement == "experience.wehelp.bullet.1"
    assert risk.export_allowed is True


def test_export_boundary_uses_safe_fallback_for_blocked_bullet_claim() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = plan.bullets[0]
    safe_text = next(
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.paragraph_id == bullet.paragraph_id
    )
    bullet.text = f"{safe_text} and documented an unsupported workflow."
    bullet.shorter_text = f"{safe_text} Unsupported workflow."
    bullet.claim_risks.append(
        ClaimRisk(
            claim="Documented an unsupported workflow",
            target_requirement="Technical documentation",
            strength=EvidenceStrength.unsupported,
            risk_level=RiskLevel.high,
            explanation="The source does not establish documentation.",
            selected_placement=bullet.paragraph_id,
            export_allowed=False,
        )
    )

    _enforce_export_boundaries(plan, document)

    assert bullet.text == safe_text
    assert bullet.shorter_text == safe_text
    assert bullet.claim_risks[0].export_allowed is False


def test_export_boundary_removes_unsafe_summary_and_global_bullet_fallbacks() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    source_by_id = {paragraph.paragraph_id: paragraph.text for paragraph in document.paragraphs}
    bullet = plan.bullets[0]
    for placement, proposal in (("summary", plan.summary), (bullet.paragraph_id, bullet)):
        proposal.text += " Administered Foundry applications."
        proposal.shorter_text = proposal.text
        plan.claim_risks.append(
            ClaimRisk(
                claim="Foundry",
                target_requirement="Foundry",
                strength=EvidenceStrength.unsupported,
                risk_level=RiskLevel.high,
                explanation="No candidate evidence for this platform.",
                selected_placement=placement,
                export_allowed=False,
            )
        )

    _enforce_export_boundaries(plan, document)

    assert plan.summary.text == plan.summary.shorter_text == source_by_id["summary"]
    assert bullet.text == bullet.shorter_text == source_by_id[bullet.paragraph_id]
    profile = TargetRoleProfile(
        title="Application Support Engineer",
        normalized_role_family="application_support",
        professional_identity="Application Support Engineer",
    )
    result = validate_rewrite_plan(document, plan, [], profile, forbidden_terms=["Foundry"])
    assert not any(issue.code.startswith("invalid_shorter") for issue in result.issues)


def test_important_supported_term_can_displace_lower_value_base_skill() -> None:
    job = parse_linkedin_simplify(Path("data/fixtures/hanmi_full.txt").read_text(encoding="utf-8"))
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(),
        rewrite_plan=identity_plan(document),
    )

    _fuse_skill_inventory(reasoning.rewrite_plan, document, keywords)
    _ensure_important_keyword_placement(reasoning, keywords, document)

    cloud = next(
        line
        for line in reasoning.rewrite_plan.skills.lines
        if line.paragraph_id == "skills.cloud_systems"
    )
    base_cloud = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.paragraph_id == "skills.cloud_systems"
    )
    base_skills = split_skill_values(base_cloud.text.split(":", 1)[1])
    assert "Operating Systems" in cloud.skills
    assert "CI/CD (GitHub Actions)" not in cloud.skills
    assert len(cloud.skills) == len(base_skills)
    assert cloud.shorter_skills == cloud.skills


def test_generic_environment_terms_do_not_displace_concrete_support_skills() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/liquid_iv_d365_technical_analyst_2026_08_12.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    compliance = next(item for item in keywords if item.normalized == "compliance")
    compliance.hiring_importance = 25
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
    tools.skills.append("Compliance")

    _fuse_skill_inventory(plan, document, keywords)

    technical = next(
        line for line in plan.skills.lines if line.paragraph_id == "skills.technical_support"
    )
    assert "Compliance" not in technical.skills
    assert "Compliance" not in tools.skills
    assert "Production Support" in technical.skills


def test_existing_base_skill_is_not_duplicated_into_another_category() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/trade_desk_support_engineer_full.txt").read_text(encoding="utf-8")
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
    apis = next(line for line in plan.skills.lines if line.paragraph_id == "skills.apis_identity")
    tools.skills.append("Postman")

    _fuse_skill_inventory(plan, document, keywords)

    assert "Postman" in apis.skills
    assert "Postman" not in tools.skills


def test_confirmed_support_vocabulary_gets_natural_experience_placement() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    terms = {"end-user support", "customer service", "service desk"}
    job = parse_linkedin_simplify(
        Path("data/fixtures/trade_desk_support_engineer_full.txt").read_text(encoding="utf-8")
    )
    keywords = grade_job_keywords(job)
    by_term = {item.normalized: item for item in keywords}

    for term in terms:
        keyword = by_term.get(term)
        if keyword is None:
            from aiadapply_v2.schemas import JobKeyword, KeywordKind, KeywordPriority

            keyword = JobKeyword(
                term=term,
                normalized=term,
                kind=KeywordKind.action,
                priority=KeywordPriority.low,
                hiring_importance=25,
                placement_utility=35,
                accepted=True,
            )
        _place_direct_category_keyword(plan, keyword)

    bullets = {item.paragraph_id: item.text for item in plan.bullets}
    assert "end-user support" in bullets["experience.csulb.bullet.1"]
    assert "customer service" in bullets["experience.wehelp.bullet.2"]
    assert "service desk" in bullets["experience.original_insurance.bullet.1"]


def test_skill_display_is_professional_and_preserves_product_casing() -> None:
    assert _display_skill("technical support") == "Technical Support"
    assert _display_skill("cross-functional collaboration") == "Cross-Functional Collaboration"
    assert _display_skill("ai") == "AI"
    assert _display_skill("PowerShell") == "PowerShell"
    assert _display_skill("OAuth 2.0") == "OAuth 2.0"


def test_supported_profile_skill_can_replace_lower_value_base_skill() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    react = JobKeyword(
        term="React",
        normalized="react",
        kind=KeywordKind.system,
        priority=KeywordPriority.high,
        occurrences=4,
        hiring_importance=92,
        placement_utility=88,
    )
    reasoning = ReasoningResult(
        role_profile=TargetRoleProfile(
            title="Frontend Support Engineer",
            normalized_role_family="application_support_engineering",
            professional_identity="Application Support Engineer",
        ),
        transferability_map=TransferabilityMap(direct_terms=["React"]),
        rewrite_plan=plan,
    )

    _fuse_skill_inventory(plan, document, [react])
    _ensure_important_keyword_placement(reasoning, [react], document)

    languages = next(
        line for line in plan.skills.lines if line.paragraph_id == "skills.languages_dbs"
    )
    assert "React" in languages.skills
    assert len(languages.skills) == len(
        split_skill_values(
            next(
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.paragraph_id == "skills.languages_dbs"
            ).split(":", 1)[1]
        )
    )


def test_rewrite_validation_rejects_skill_without_candidate_evidence() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/trade_desk_support_engineer_full.txt").read_text(encoding="utf-8")
    )
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
    tools.skills[-1] = "Rust"

    validation = validate_rewrite_plan(
        document,
        plan,
        [],
        build_target_role_profile(job, grade_job_keywords(job)),
    )

    assert not validation.passed
    assert any(issue.code == "unsupported_skill_inserted" for issue in validation.issues)


def test_skill_fusion_prioritizes_supported_job_tools_and_drops_weak_panel_noise() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/anduril_rotation_full.txt").read_text(encoding="utf-8")
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
    tools.skills = ["Compliance", "Project Management", "AI", "Computer Vision"]
    tools.shorter_skills = list(tools.skills)

    _fuse_skill_inventory(plan, document, keywords)

    assert "Project Management" in tools.skills
    assert len(tools.skills) <= 7
    assert "Compliance" not in tools.skills
    assert "AI" not in tools.skills
    assert "Computer Vision" not in tools.skills


def test_skill_fusion_drops_unverified_additions_and_keeps_base_inventory_size() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/pds_health_epic_analyst_full_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
    technical = next(
        line for line in plan.skills.lines if line.paragraph_id == "skills.technical_support"
    )
    tools.skills = ["Jira", "Confluence", "Excel", "Outlook"]
    tools.shorter_skills = list(tools.skills)
    technical.skills = ["Production Support", "Troubleshooting", "Customer Service"]
    technical.shorter_skills = list(technical.skills)

    _fuse_skill_inventory(plan, document, keywords)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(),
        rewrite_plan=plan,
    )
    supported = [
        keyword
        for keyword in keywords
        if keyword.normalized
        not in {"epic", "epic certification", "professional billing", "claims"}
    ]
    _ensure_important_keyword_placement(reasoning, supported, document)

    assert "Outlook" not in tools.skills
    assert "VS Code" in tools.skills
    base_technical = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.paragraph_id == "skills.technical_support"
    )
    base_skills = [value.strip() for value in base_technical.text.split(":", 1)[1].split(",")]
    assert 1 <= len(technical.skills) <= len(base_skills)
    assert "Troubleshooting" in technical.skills


def test_skill_budget_uses_measured_spare_width_without_changing_the_fallback_inventory() -> None:
    document = parse_resume_docx(BASE)
    skill_paragraphs = [
        paragraph for paragraph in document.paragraphs if paragraph.kind.value == "skill_line"
    ]
    for paragraph in skill_paragraphs:
        paragraph.rendered_max_width_points = 560.0
    tools = next(
        paragraph for paragraph in skill_paragraphs if paragraph.paragraph_id == "skills.tools"
    )
    tools.rendered_max_width_points = 315.0

    budget = _measured_skill_character_budget(document, tools)

    assert budget > tools.character_budget
    assert budget <= round(tools.character_budget * 1.5)


def test_skill_row_overflow_drops_low_value_tail_without_adding_a_row() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    apis = next(line for line in plan.skills.lines if line.paragraph_id == "skills.apis_identity")
    apis.skills = [
        "Microsoft 365",
        "Entra ID",
        "MFA",
        "RBAC",
        "SSO",
        "Microsoft Graph API",
        "REST APIs",
        "Webhooks",
        "OAuth 2.0",
        "SAML",
        "Postman",
    ]
    apis.shorter_skills = list(apis.skills)
    changed = _repair_skill_row_overflow(plan, document, [], {apis.paragraph_id: 2})
    assert changed
    assert "Postman" not in apis.skills
    assert len(apis.skills) == 10


def test_wrapped_skill_row_is_repaired_even_when_baseline_also_wraps() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    apis = next(line for line in plan.skills.lines if line.paragraph_id == "skills.apis_identity")
    original = list(apis.skills)

    changed = _repair_skill_row_overflow(
        plan,
        document,
        [],
        {apis.paragraph_id: 2},
        {apis.paragraph_id: 2},
    )

    assert changed
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
    assert "Postman" not in apis.skills
    assert "Postman" in tools.skills
    assert set(apis.skills + tools.skills) >= set(original)


def test_high_value_postman_can_move_to_tools_instead_of_being_dropped() -> None:
    document = parse_resume_docx(BASE)
    for paragraph in document.paragraphs:
        if paragraph.kind.value == "skill_line":
            paragraph.rendered_max_width_points = 560.0
    tools_meta = next(
        paragraph for paragraph in document.paragraphs if paragraph.paragraph_id == "skills.tools"
    )
    tools_meta.rendered_max_width_points = 400.0
    plan = identity_plan(document)
    apis = next(line for line in plan.skills.lines if line.paragraph_id == "skills.apis_identity")
    keywords = [
        JobKeyword(
            term=value,
            normalized=value.casefold(),
            priority=KeywordPriority.high,
            kind=KeywordKind.system,
            occurrences=2,
            hiring_importance=60 if value == "Postman" else 80,
            placement_utility=80,
        )
        for value in apis.skills
    ]

    changed = _repair_skill_row_overflow(plan, document, keywords, {apis.paragraph_id: 2})
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")

    assert changed
    assert "Postman" not in apis.skills
    assert "Postman" in tools.skills


def test_postman_relocation_preserves_other_skills_and_shorter_variant() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    apis = next(line for line in plan.skills.lines if line.paragraph_id == "skills.apis_identity")
    before = list(apis.skills)
    apis.shorter_skills = ["Postman", "REST APIs"]
    keywords = [
        JobKeyword(
            term="Postman",
            normalized="postman",
            priority=KeywordPriority.high,
            kind=KeywordKind.system,
            occurrences=2,
            hiring_importance=99,
            placement_utility=99,
        )
    ]
    assert _repair_skill_row_overflow(plan, document, keywords, {apis.paragraph_id: 2})
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
    assert apis.skills == [value for value in before if value != "Postman"]
    assert apis.shorter_skills == ["REST APIs"]
    assert "Postman" in tools.skills
    assert "Postman" in tools.shorter_skills


def test_postman_displaced_before_rendering_is_retained_in_tools() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    apis = next(line for line in plan.skills.lines if line.paragraph_id == "skills.apis_identity")
    apis.skills.remove("Postman")
    job = parse_linkedin_simplify(
        Path("data/fixtures/pacific_life_platform_engineer_ii_live_2026_09_21.txt").read_text()
    )
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, []),
        transferability_map=TransferabilityMap(),
        rewrite_plan=plan,
    )
    for _ in range(2):
        _ensure_important_keyword_placement(reasoning, [], document)
        tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
        assert tools.skills.count("Postman") == 1
        assert tools.shorter_skills.count("Postman") == 1


def test_equal_length_model_fallback_uses_shorter_source_without_changing_primary() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = plan.bullets[0]
    source = bullet.text
    bullet.text = source[:-1] + " and service performance."
    bullet.shorter_text = source[:-1] + " for service performance."
    primary = bullet.text
    assert len(bullet.text) == len(bullet.shorter_text)
    _repair_equal_length_fallbacks(plan, document)
    assert bullet.text == primary
    assert bullet.shorter_text == source
    bullet.text = "Diagnosed incidents."
    bullet.shorter_text = bullet.text
    _repair_equal_length_fallbacks(plan, document)
    assert bullet.shorter_text == "Diagnosed incidents."


def test_postman_does_not_bounce_back_after_destination_overflows() -> None:
    document = parse_resume_docx(BASE)
    for paragraph in document.paragraphs:
        if paragraph.kind.value == "skill_line":
            paragraph.character_budget = 500  # Simulate optimistic width estimates.
    plan = identity_plan(document)
    apis = next(line for line in plan.skills.lines if line.paragraph_id == "skills.apis_identity")
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
    keywords = [
        JobKeyword(
            term="Postman",
            normalized="postman",
            priority=KeywordPriority.high,
            kind=KeywordKind.system,
            occurrences=2,
            hiring_importance=99,
            placement_utility=99,
        )
    ]
    moved: set[str] = set()
    assert _repair_skill_row_overflow(
        plan,
        document,
        keywords,
        {apis.paragraph_id: 2, tools.paragraph_id: 1},
        relocated_skills=moved,
    )
    assert moved == {"postman"}
    assert _repair_skill_row_overflow(
        plan,
        document,
        keywords,
        {apis.paragraph_id: 1, tools.paragraph_id: 2},
        relocated_skills=moved,
    )
    assert "Postman" in tools.skills
    assert "Postman" not in apis.skills


def test_skill_repair_uses_exported_compressed_and_restored_values(tmp_path: Path) -> None:
    document = parse_resume_docx(BASE)
    for mode in ("compressed", "restored"):
        plan = identity_plan(document)
        apis = next(
            line for line in plan.skills.lines if line.paragraph_id == "skills.apis_identity"
        )
        apis.skills = ["REST APIs"]
        apis.shorter_skills = ["Microsoft Graph API", "Postman"]
        candidate = write_resume_candidate(
            document,
            plan,
            tmp_path / f"{mode}.docx",
            compressed_paragraph_ids={apis.paragraph_id} if mode == "compressed" else set(),
            reverted_paragraph_ids={apis.paragraph_id} if mode == "restored" else set(),
        )
        _use_rendered_skill_rows(plan, parse_resume_docx(candidate))
        assert "Postman" in apis.skills
        assert _repair_skill_row_overflow(plan, document, [], {apis.paragraph_id: 2})
        final = write_resume_candidate(document, plan, tmp_path / f"{mode}-repaired.docx")
        actual = {p.paragraph_id: p.text for p in parse_resume_docx(final).paragraphs}
        assert "Postman" not in actual[apis.paragraph_id]
        assert "Postman" in actual["skills.tools"]


def test_reversed_model_variants_are_ordered_before_factual_validation() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = plan.bullets[0]
    original = bullet.text
    bullet.shorter_text = original + " Extra context."
    _order_rewrite_variants(plan)
    assert bullet.text == original + " Extra context."
    assert bullet.shorter_text == original
    # Reordering is not an exemption from the source evidence checks.
    bullet.text = bullet.text.replace("Microsoft 365", "a different tool")
    _restore_unjustified_shortening(plan, document, [])
    assert bullet.text == original


def test_unjustified_or_severe_shortening_restores_source_bullet() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/pds_health_epic_analyst_full_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    paragraph_id = "experience.original_insurance.bullet.4"
    bullet = next(item for item in plan.bullets if item.paragraph_id == paragraph_id)
    source = bullet.text
    bullet.text = "Administered Microsoft 365 and Entra ID configurations."

    _restore_unjustified_shortening(plan, document, keywords)

    assert bullet.text == source


def test_rewording_without_a_new_target_term_restores_source_bullet() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    paragraph_id = "experience.csulb.bullet.2"
    bullet = next(item for item in plan.bullets if item.paragraph_id == paragraph_id)
    source = bullet.text
    bullet.text = source.replace("resetting user access", "restoring user access")

    _restore_unjustified_shortening(plan, document, [])

    assert bullet.text == source


def test_rewrite_that_drops_a_named_source_system_is_restored() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/pds_health_epic_analyst_full_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    paragraph_id = "experience.original_insurance.bullet.3"
    bullet = next(item for item in plan.bullets if item.paragraph_id == paragraph_id)
    source = bullet.text
    bullet.text = source.replace("production database", "production") + " Improved efficiency."

    _restore_unjustified_shortening(plan, document, keywords)

    assert bullet.text == source


def test_resolved_nonexportable_gap_is_removed_from_review_flags() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/pds_health_epic_analyst_full_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    plan = identity_plan(parse_resume_docx(BASE))
    plan.claim_risks.append(
        ClaimRisk(
            claim="Report creation",
            target_requirement="report creation",
            strength=EvidenceStrength.weakly_transferable,
            risk_level=RiskLevel.medium,
            explanation="Excluded from export.",
            selected_placement="keyword gap only",
            export_allowed=False,
        )
    )

    _prune_resolved_nonexportable_risks(plan, keywords)

    assert plan.claim_risks == []


def test_absent_transferable_claim_is_removed_from_review_flags() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/pds_health_epic_analyst_full_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    plan = identity_plan(parse_resume_docx(BASE))
    plan.claim_risks.append(
        ClaimRisk(
            claim="System maintenance",
            target_requirement="system maintenance",
            strength=EvidenceStrength.strongly_transferable,
            risk_level=RiskLevel.low,
            explanation="Review this adjacent duty.",
            selected_placement="projects.loavenly.bullet.1",
            export_allowed=True,
        )
    )

    _prune_absent_claim_risks(plan, keywords)

    assert plan.claim_risks == []


def test_rewrite_validation_allows_removal_of_a_verified_base_skill() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/pds_health_epic_analyst_full_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    technical = next(
        line for line in plan.skills.lines if line.paragraph_id == "skills.technical_support"
    )
    technical.skills.remove("Root Cause Analysis")

    validation = validate_rewrite_plan(
        document,
        plan,
        [],
        build_target_role_profile(job, grade_job_keywords(job)),
    )

    assert validation.passed
    assert not any(issue.code == "base_skill_removed" for issue in validation.issues)


def test_layout_fallbacks_preserve_source_evidence_instead_of_overcompressing() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    plan.summary.text = f"{plan.summary.text} Strengths include customer service."
    plan.summary.shorter_text = "Application Support Engineer with customer service."
    volunteer = next(
        item for item in plan.bullets if item.paragraph_id == "experience.wehelp.bullet.2"
    )
    volunteer.shorter_text = "Supported users and fixed issues."

    _sanitize_shorter_fallbacks(plan, document)

    source_by_id = {item.paragraph_id: item.text for item in document.paragraphs}
    assert plan.summary.shorter_text == source_by_id["summary"]
    assert volunteer.shorter_text == source_by_id["experience.wehelp.bullet.2"]


def test_draft_project_fallback_can_fit_without_erasing_new_technology() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    project = next(
        item for item in plan.bullets if item.paragraph_id == "projects.loavenly.bullet.2"
    )
    compact = (
        "Configured RBAC and Java API; fixed authentication, synchronization, "
        "and permission issues in production."
    )
    assert len(compact) < 0.90 * len(project.text)
    assert len(compact) >= 0.75 * len(project.text)
    project.text = compact
    project.shorter_text = compact

    _sanitize_shorter_fallbacks(plan, document, draft_terms={"java"})

    assert project.shorter_text == compact


def test_unproven_issue_reproduction_is_kept_as_investigation() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    insurance = next(
        item
        for item in plan.bullets
        if item.paragraph_id == "experience.original_insurance.bullet.2"
    )
    insurance.text = "Reproduced and resolved production issues."
    insurance.shorter_text = "Reproducing production issues."

    _remove_unproven_reproduction(plan, document)

    assert insurance.text == "Investigated and resolved production issues."
    assert insurance.shorter_text == "Investigating production issues."


def test_primary_summary_cannot_drop_a_named_source_system() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    plan.summary.text = plan.summary.text.replace(
        " and built Loavenly, a production multi-tenant SaaS platform",
        "",
    )

    _restore_unjustified_shortening(plan, document, [])

    source = next(item.text for item in document.paragraphs if item.paragraph_id == "summary")
    assert plan.summary.text == source


def test_role_identity_is_reapplied_after_summary_source_restoration() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/trade_desk_platform_support_analyst_i_2026_08_06.txt").read_text(
            encoding="utf-8"
        )
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(),
        rewrite_plan=identity_plan(document),
    )
    reasoning.rewrite_plan.summary.text = "Platform Support Analyst with production support."

    _restore_unjustified_shortening(reasoning.rewrite_plan, document, keywords)
    _ensure_target_identity(reasoning)
    _sanitize_shorter_fallbacks(reasoning.rewrite_plan, document)
    _ensure_target_identity(reasoning)

    assert reasoning.rewrite_plan.summary.text.startswith("Application Support Analyst")
    assert reasoning.rewrite_plan.summary.shorter_text.startswith("Application Support Analyst")
    assert "Loavenly" in reasoning.rewrite_plan.summary.text


def test_compression_prioritizes_measured_overflow_paragraph() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    overflow = "experience.original_insurance.bullet.3"
    layout = LayoutResult(
        passed=False,
        page_count=2,
        rendered_lines=62,
        overflow_paragraph_ids=[overflow, "PROJECTS"],
    )

    selected = _select_compression_candidate(document, plan, layout, set())

    assert selected == overflow


def test_expansion_restores_shortened_paragraph_before_shifted_anchor() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    paragraph_id = "experience.original_insurance.bullet.3"
    bullet = next(item for item in plan.bullets if item.paragraph_id == paragraph_id)
    bullet.text = "Monitored ETL jobs."
    layout = LayoutResult(
        passed=False,
        page_count=1,
        rendered_lines=50,
        section_anchor_deltas={"PROJECTS": -12.0},
        paragraph_line_counts={paragraph_id: 1},
        baseline_paragraph_line_counts={paragraph_id: 2},
    )

    selected = _select_expansion_candidate(document, plan, layout, set())

    assert selected == paragraph_id


def test_overflow_reversion_restores_only_the_changed_paragraph_that_still_overflows() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    paragraph_id = "projects.loavenly.bullet.2"
    bullet = next(item for item in plan.bullets if item.paragraph_id == paragraph_id)
    bullet.text = f"{bullet.text} with additional platform configuration detail"
    layout = LayoutResult(
        passed=False,
        page_count=1,
        rendered_lines=61,
        overflow_paragraph_ids=[paragraph_id, "bottom-edge"],
        paragraph_line_counts={paragraph_id: 3},
        baseline_paragraph_line_counts={paragraph_id: 2},
    )

    selected = _select_overflow_reversion_candidate(document, plan, layout, set())

    assert selected == paragraph_id


def test_overflow_reversion_uses_shifted_section_when_bottom_edge_is_only_signal() -> None:
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    paragraph_id = "projects.loavenly.bullet.1"
    bullet = next(item for item in plan.bullets if item.paragraph_id == paragraph_id)
    bullet.text = f"{bullet.text} with additional workflow automation detail"
    layout = LayoutResult(
        passed=False,
        page_count=1,
        rendered_lines=61,
        overflow_paragraph_ids=["bottom-edge"],
        section_anchor_deltas={"EDUCATION": 12.8},
        paragraph_line_counts={paragraph_id: 3},
        baseline_paragraph_line_counts={paragraph_id: 2},
    )

    selected = _select_overflow_reversion_candidate(document, plan, layout, set())

    assert selected == paragraph_id


def test_dense_erp_keywords_are_balanced_across_fixed_skill_rows() -> None:
    job = parse_linkedin_simplify(
        Path("data/fixtures/liquid_iv_full.txt").read_text(encoding="utf-8")
    )
    keywords = grade_job_keywords(job)
    document = parse_resume_docx(BASE)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(),
        rewrite_plan=identity_plan(document),
    )

    _ensure_important_keyword_placement(reasoning, keywords, document)

    by_id = {paragraph.paragraph_id: paragraph for paragraph in document.paragraphs}
    for line in reasoning.rewrite_plan.skills.lines:
        paragraph = by_id[line.paragraph_id]
        _label, values = paragraph.text.split(":", 1)
        base_items = split_skill_values(values)
        assert len(line.skills) <= len(base_items)
        assert line.skills


def test_added_keywords_use_final_bullet_text_not_model_annotations():
    from aiadapply_v2.pipeline import _build_change_manifest
    from aiadapply_v2.schemas import ResumeDocument

    base = parse_resume_docx(BASE)
    plan = identity_plan(base)
    final = ResumeDocument.model_validate(base.model_dump())
    key = "experience.original_insurance.bullet.3"
    paragraph = next(p for p in final.paragraphs if p.paragraph_id == key)
    paragraph.text = paragraph.text.replace(
        "a Python reconciliation script", "Python reconciliation automation"
    )
    bullet = next(b for b in plan.bullets if b.paragraph_id == key)
    bullet.target_terms = ["Kubernetes", "SQL"]
    keywords = grade_job_keywords(
        parse_linkedin_simplify(
            Path("data/fixtures/mendix_application_support_engineer_2026_09_22.txt")
            .read_text()
            .replace(
                "What You'll Be Doing",
                "What You'll Be Doing\nBuild Python automation and monitor ETL jobs using SQL and Kubernetes.",
            )
        )
    )
    changes = _build_change_manifest(base, final, plan, set(), keywords)
    change = next(c for c in changes if c.paragraph_id == key)
    assert change.added_terms == ["automation"]
    assert all(not c.added_terms for c in changes if c.paragraph_id != key)


def test_action_inflection_without_new_requirement_coverage_is_restored():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = next(b for b in plan.bullets if b.paragraph_id == "projects.loavenly.bullet.2")
    bullet.text = bullet.text.replace("resolved", "troubleshot")
    bullet.target_terms = ["troubleshooting"]
    expected = next(p.text for p in document.paragraphs if p.paragraph_id == bullet.paragraph_id)
    job = parse_linkedin_simplify(
        Path("data/fixtures/mendix_application_support_engineer_2026_09_22.txt").read_text()
    )
    _restore_unjustified_shortening(plan, document, grade_job_keywords(job))
    assert bullet.text == expected


def test_soft_skill_ranking_cannot_displace_technical_skills_or_sla_management():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    line = next(
        line for line in plan.skills.lines if line.paragraph_id == "skills.technical_support"
    )
    line.skills = [
        "Problem-Solving",
        "Troubleshooting",
        "Incident Response",
        "Root Cause Analysis",
        "Service Level Agreements",
    ]
    job = parse_linkedin_simplify(
        Path("data/fixtures/mendix_application_support_engineer_2026_09_22.txt").read_text()
    )
    keywords = grade_job_keywords(job)
    _fuse_skill_inventory(plan, document, keywords)
    reasoning = ReasoningResult(
        role_profile=build_target_role_profile(job, keywords),
        transferability_map=TransferabilityMap(),
        rewrite_plan=plan,
    )
    _ensure_important_keyword_placement(
        reasoning, [k for k in keywords if k.normalized == "problem-solving"], document
    )
    assert "Problem-Solving" not in line.skills
    assert "SLA Management" in line.skills
    assert "Incident Response" in line.skills


def test_rearranged_existing_claim_without_new_role_value_is_restored():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = next(b for b in plan.bullets if b.paragraph_id == "experience.original_insurance.bullet.1")
    bullet.text = "Handled agency management SaaS, internal integrations, and Microsoft 365 incidents through Jira; resolved 200+ production support tickets while maintaining sub-90-minute average resolution."
    bullet.target_terms = []
    expected = next(p.text for p in document.paragraphs if p.paragraph_id == bullet.paragraph_id)
    job = parse_linkedin_simplify(Path("data/fixtures/mendix_application_support_engineer_2026_09_22.txt").read_text())
    feedback = _restore_unjustified_shortening(plan, document, grade_job_keywords(job))
    assert bullet.text == expected
    assert any(bullet.paragraph_id in message for message in feedback)


def test_restoration_explains_lost_context_to_correction_model():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = next(b for b in plan.bullets if b.paragraph_id == "experience.original_insurance.bullet.4")
    original = bullet.text
    bullet.text = original.replace("license management", "other administration")
    feedback = _restore_unjustified_shortening(plan, document, [])
    assert bullet.text == original
    assert any(bullet.paragraph_id in message and "license management" in message for message in feedback)


def test_rest_api_project_rewrite_keeps_generic_api_evidence():
    from aiadapply_v2.text import preserves_source_term, contains_term
    assert preserves_source_term("Configured REST APIs", "API")
    assert not contains_term("Configured REST APIs", "API")
    assert not preserves_source_term("Configured an API", "Microsoft Graph API")
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = next(b for b in plan.bullets if b.paragraph_id == "projects.loavenly.bullet.2")
    bullet.text = "Configured RBAC; fixed REST APIs, Webhooks, authentication, synchronization, and permission issues in live production operations."
    bullet.shorter_text = bullet.text
    expected = bullet.text
    feedback = _restore_unjustified_shortening(plan, document, [], draft_terms={"rest apis", "webhooks"})
    assert bullet.text == expected
    assert not any("dropped protected phrases: API" in item for item in feedback)
    _sanitize_shorter_fallbacks(plan, document, draft_terms={"rest apis", "webhooks"})
    assert bullet.shorter_text == expected


def test_floqast_cosmetic_reframes_are_restored_in_both_variants():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    examples = {
        "experience.csulb.bullet.3": "Coordinated with campus IT and third-party vendors on incidents, maintaining documentation, escalation notes, and tracking of recurring issues.",
        "experience.wehelp.bullet.2": "Keep Loavenly ready for live Wednesday and Saturday distributions through production troubleshooting, user access management, update validation, and training staff and volunteers.",
        "projects.loavenly.bullet.1": "Enabled client intake, inventory tracking, reporting, and RBAC across three food bank locations by building, deploying, and operating a production multi-tenant SaaS platform.",
    }
    source = {p.paragraph_id: p.text for p in document.paragraphs}
    for bullet in plan.bullets:
        if bullet.paragraph_id in examples:
            bullet.text = examples[bullet.paragraph_id]
            bullet.shorter_text = examples[bullet.paragraph_id]
            bullet.target_terms = ["IT", "documentation", "troubleshooting", "SaaS"]
    job = parse_linkedin_simplify(Path("data/fixtures/floqast_full.txt").read_text())
    _restore_unjustified_shortening(plan, document, grade_job_keywords(job))
    for bullet in plan.bullets:
        if bullet.paragraph_id in examples:
            assert bullet.text == source[bullet.paragraph_id]
            assert bullet.shorter_text == source[bullet.paragraph_id]


def test_skill_ranking_preserves_python_and_troubleshooting_source_order():
    from aiadapply_v2.pipeline import _rank_skill_lines
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    original = {line.paragraph_id: list(line.skills) for line in plan.skills.lines}
    for line in plan.skills.lines:
        line.skills.reverse()
    job = parse_linkedin_simplify(Path("data/fixtures/floqast_full.txt").read_text())
    keywords = grade_job_keywords(job)
    _fuse_skill_inventory(plan, document, keywords)
    _rank_skill_lines(plan, document, keywords)
    for line in plan.skills.lines:
        assert line.skills == original[line.paragraph_id]


def test_supported_ticketing_term_survives_without_rewrite_quota():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = next(b for b in plan.bullets if b.paragraph_id == "experience.original_insurance.bullet.1")
    bullet.text = bullet.text.replace("via Jira", "through Jira ticketing")
    expected = bullet.text
    job = parse_linkedin_simplify(Path("data/fixtures/floqast_full.txt").read_text())
    _restore_unjustified_shortening(plan, document, grade_job_keywords(job))
    assert bullet.text == expected


def test_inflection_only_keyword_does_not_license_full_bullet_rearrangement():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    job = parse_linkedin_simplify(Path("data/fixtures/floqast_full.txt").read_text())
    keywords = grade_job_keywords(job)
    bullet = next(b for b in plan.bullets if b.paragraph_id == "experience.csulb.bullet.2")
    original = bullet.text
    bullet.text = "Reproduced access and scheduling errors to diagnose configuration and permission issues, reset user access, and escalated complex incidents to campus IT for further investigation."
    bullet.shorter_text = bullet.text
    _restore_unjustified_shortening(plan, document, keywords)
    assert bullet.text == original.replace("configurations", "configuration")
    assert bullet.shorter_text == bullet.text
    project = next(b for b in plan.bullets if b.paragraph_id == "projects.loavenly.bullet.2")
    original_project = project.text
    project.text = "Resolved live production API, authentication, synchronization, and permission issues alongside hands-on RBAC configuration."
    project.shorter_text = project.text
    _restore_unjustified_shortening(plan, document, keywords)
    assert project.text == original_project
    assert project.shorter_text == original_project


def test_cosmetic_shorter_fallback_cannot_bypass_meaningful_edit_guard():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = next(b for b in plan.bullets if b.paragraph_id == "experience.original_insurance.bullet.1")
    original = bullet.text
    bullet.shorter_text = original.replace("via Jira", "in Jira")
    _restore_unjustified_shortening(plan, document, [])
    assert bullet.text == original
    assert bullet.shorter_text == original


def test_soft_skill_appendix_is_restored_without_blocking_real_ticket_context():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = next(b for b in plan.bullets if b.paragraph_id == "experience.original_insurance.bullet.2")
    original = bullet.text
    bullet.text = original.rstrip(".") + ", demonstrating critical thinking."
    bullet.shorter_text = bullet.text
    job = parse_linkedin_simplify(Path("data/fixtures/floqast_full.txt").read_text())
    _restore_unjustified_shortening(plan, document, grade_job_keywords(job))
    assert bullet.text == original
    assert bullet.shorter_text == original


def _scoped_risk_fixture(term, placement, evidence_id, strength=EvidenceStrength.direct):
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    risk = ClaimRisk(claim=term, target_requirement=term, strength=EvidenceStrength.weakly_transferable,
                     risk_level=RiskLevel.medium, explanation="Check source scope", selected_placement=placement, export_allowed=False)
    plan.claim_risks.append(risk)
    reasoning = ReasoningResult(
        role_profile=TargetRoleProfile(title="Support", normalized_role_family="technical_support_integrations", professional_identity="Technical Support Engineer"),
        transferability_map=TransferabilityMap(matches=[EvidenceMatch(
            target_term=term, target_requirement=term, evidence_id=evidence_id,
            source_text="", semantic_score=1, action_compatibility=1, system_compatibility=1,
            environment_compatibility=1, outcome_compatibility=1, strength=strength,
            reasoning="fixture", suggested_placement=placement)]),
        rewrite_plan=plan)
    return document, reasoning


def test_global_profile_does_not_clear_employer_knowledge_base_risk():
    _, reasoning = _scoped_risk_fixture("knowledge base", "experience.csulb.bullet.3", "evidence.candidate_profile.knowledge_base")
    _prune_direct_evidence_risks(reasoning)
    assert len(reasoning.rewrite_plan.claim_risks) == 1


def test_actual_ticket_context_clears_unnecessary_case_management_question():
    from aiadapply_v2.evidence.loaders import build_evidence_graph
    document, reasoning = _scoped_risk_fixture("case management", "experience.original_insurance.bullet.1", "evidence.experience.original_insurance.bullet.1", EvidenceStrength.strongly_transferable)
    _prune_direct_evidence_risks(reasoning, build_evidence_graph(document, []))
    assert reasoning.rewrite_plan.claim_risks == []


def test_restored_fallback_longer_than_meaningful_primary_reuses_primary():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    bullet = next(b for b in plan.bullets if b.paragraph_id == "experience.original_insurance.bullet.1")
    source = bullet.text
    primary = source.replace("tickets via Jira", "cases via Jira ticketing").replace(", maintaining", ";")
    bullet.text = primary
    bullet.shorter_text = source.replace("via Jira", "in Jira")
    job = parse_linkedin_simplify(Path("data/fixtures/floqast_full.txt").read_text())
    keywords = grade_job_keywords(job)
    assert len(primary) < len(source)
    _restore_unjustified_shortening(plan, document, keywords)
    assert bullet.text == primary
    assert bullet.shorter_text == primary
    result = validate_rewrite_plan(document, plan, keywords, build_target_role_profile(job, keywords))
    assert not any(issue.code == "invalid_shorter_candidate" and issue.paragraph_id == bullet.paragraph_id for issue in result.issues)


def test_unchanged_skills_manifest_does_not_claim_reordering():
    from aiadapply_v2.pipeline import _build_change_manifest
    document = parse_resume_docx(BASE)
    changes = _build_change_manifest(document, document, identity_plan(document), set())
    skills = [change for change in changes if change.paragraph_id.startswith("skills.")]
    assert len(skills) == 5
    assert all(change.change_type == "unchanged" for change in skills)
    assert all(change.explanation == "Kept the original skills and their order." for change in skills)


def test_confirmed_generic_ai_survives_skills_normalization_and_layout_selection():
    document = parse_resume_docx(BASE)
    plan = identity_plan(document)
    ai = JobKeyword(term="AI", normalized="ai", kind="environment", hiring_importance=55, placement_utility=55)
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
    tools.skills.append("AI")
    tools.shorter_skills.append("AI")
    _fuse_skill_inventory(plan, document, [ai], unsupported_terms=[])
    assert "AI" in tools.skills and "AI" in tools.shorter_skills
    reasoning = ReasoningResult(role_profile=TargetRoleProfile(title="Analyst", normalized_role_family="product_operations", professional_identity="Application Support Specialist"), transferability_map=TransferabilityMap(), rewrite_plan=plan)
    _ensure_important_keyword_placement(reasoning, [ai], document)
    assert "AI" in tools.skills and "AI" in tools.shorter_skills
    assert all("AI" not in line.skills for line in plan.skills.lines if line.paragraph_id != "skills.tools")


def test_ai_skill_allowance_does_not_allow_unsupported_ai_or_named_tool_bundles():
    document = parse_resume_docx(BASE)
    ai = JobKeyword(term="AI", normalized="ai", kind="environment", hiring_importance=55, placement_utility=55)
    for value, unsupported in (("AI", ["AI"]), ("AI (Qualtrics, Zendesk)", [])):
        plan = identity_plan(document)
        tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
        tools.skills.append(value)
        _fuse_skill_inventory(plan, document, [ai], unsupported_terms=unsupported)
        assert not any("AI" in item or "Qualtrics" in item or "Zendesk" in item for line in plan.skills.lines for item in line.skills)
    ai.accepted = False
    plan = identity_plan(document)
    tools = next(line for line in plan.skills.lines if line.paragraph_id == "skills.tools")
    tools.skills.append("AI")
    _fuse_skill_inventory(plan, document, [ai])
    assert "AI" not in tools.skills
