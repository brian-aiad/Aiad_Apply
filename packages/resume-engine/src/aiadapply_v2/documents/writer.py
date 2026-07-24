from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from aiadapply_v2.schemas import ResumeDocument, RewritePlan


def write_resume_candidate(
    base: ResumeDocument,
    plan: RewritePlan,
    output_path: str | Path,
    *,
    compression_level: int = 0,
    compressed_paragraph_ids: set[str] | None = None,
) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document(str(base.source_path))
    by_id = {paragraph.paragraph_id: paragraph for paragraph in base.paragraphs}

    summary_meta = by_id["summary"]
    summary_text = _select_text(
        plan.summary.text,
        plan.summary.shorter_text,
        summary_meta.character_budget,
        compression_level,
        force_shorter=(
            compressed_paragraph_ids is not None and "summary" in compressed_paragraph_ids
        ),
    )
    _replace_all_text(document.paragraphs[summary_meta.index], summary_text)

    proposed_skills = {line.paragraph_id: line for line in plan.skills.lines}
    for paragraph_id, proposed_skill in proposed_skills.items():
        meta = by_id[paragraph_id]
        selected_skills = (
            proposed_skill.shorter_skills
            if (compressed_paragraph_ids is not None and paragraph_id in compressed_paragraph_ids)
            or compression_level >= 2
            or (
                compressed_paragraph_ids is None
                and compression_level >= 1
                and len(", ".join(proposed_skill.skills)) > meta.character_budget
            )
            else proposed_skill.skills
        )
        _replace_skill_text(
            document.paragraphs[meta.index],
            proposed_skill.category,
            ", ".join(selected_skills),
        )

    proposed_bullets = {bullet.paragraph_id: bullet for bullet in plan.bullets}
    for paragraph_id, proposed_bullet in proposed_bullets.items():
        meta = by_id[paragraph_id]
        selected_text = _select_text(
            proposed_bullet.text,
            proposed_bullet.shorter_text,
            meta.character_budget,
            compression_level,
            force_shorter=(
                compressed_paragraph_ids is not None and paragraph_id in compressed_paragraph_ids
            ),
        )
        _replace_all_text(document.paragraphs[meta.index], selected_text)

    document.save(str(output))
    return output


def _select_text(
    full: str,
    shorter: str,
    budget: int,
    compression_level: int,
    *,
    force_shorter: bool = False,
) -> str:
    if force_shorter:
        return shorter
    if compression_level >= 2:
        return shorter
    if compression_level >= 1 and len(full) > budget:
        return shorter
    return full


def _replace_all_text(paragraph: Paragraph, value: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(value)
        return
    runs[0].text = value
    for run in runs[1:]:
        run.text = ""


def _replace_skill_text(paragraph: Paragraph, category: str, skills: str) -> None:
    existing = paragraph.text
    prefix_match = re.match(r"^.*?:\s*", existing)
    prefix = prefix_match.group(0) if prefix_match else f"{category}: "
    runs = paragraph.runs
    if not runs:
        label = paragraph.add_run(prefix)
        label.bold = True
        paragraph.add_run(skills)
        return
    runs[0].text = prefix
    runs[0].bold = True
    if len(runs) == 1:
        content = paragraph.add_run(skills)
        content.bold = False
        return
    runs[1].text = skills
    runs[1].bold = False
    for run in runs[2:]:
        run.text = ""
