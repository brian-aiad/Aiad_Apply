from __future__ import annotations

import hashlib
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.text.hyperlink import Hyperlink

from aiadapply_v2.schemas import (
    ParagraphKind,
    ResumeDocument,
    ResumeParagraph,
    ResumeSection,
    TextRun,
)
from aiadapply_v2.text import slug

SECTION_HEADINGS = {"SKILLS", "EXPERIENCE", "PROJECTS", "EDUCATION", "CERTIFICATIONS"}
ENTRY_IDENTITIES = {
    "ORIGINAL INSURANCE GROUP INC": "original_insurance",
    "CALIFORNIA STATE UNIVERSITY, LONG BEACH": "csulb",
    "WEHELP, CHRIST THE GOOD SHEPHERD FOOD BANK": "wehelp",
    "LOAVENLY": "loavenly",
}
EDITABLE_SECTIONS = {
    "summary",
    "skills",
    "experience.original_insurance",
    "experience.csulb",
    "experience.wehelp",
    "projects.loavenly",
}
METRIC_PATTERNS = (
    re.compile(r"\b\d+\+"),
    re.compile(r"\bsub-\d+(?:-minute|ms)\b", re.IGNORECASE),
    re.compile(
        r"\b\d+\s+(?:carrier/vendor feeds|hours?|minutes?|SaaS apps?|food bank locations?)\b",
        re.IGNORECASE,
    ),
    re.compile(r"\bthree\s+(?:distribution|food bank)\s+locations\b", re.IGNORECASE),
)


def parse_resume_docx(path: str | Path) -> ResumeDocument:
    source = Path(path).resolve()
    raw = source.read_bytes()
    document = Document(str(source))
    paragraphs = _parse_paragraphs(document)
    sections = _build_sections(paragraphs)
    protected_strings = _protected_strings(paragraphs)
    hyperlinks = _hyperlinks(document)
    metrics = _metrics_by_section(paragraphs)
    section = document.sections[0]
    margins = {
        "top": _points(section.top_margin),
        "right": _points(section.right_margin),
        "bottom": _points(section.bottom_margin),
        "left": _points(section.left_margin),
    }
    return ResumeDocument(
        source_path=source,
        source_sha256=hashlib.sha256(raw).hexdigest(),
        paragraphs=paragraphs,
        sections=sections,
        protected_strings=protected_strings,
        protected_hyperlinks=hyperlinks,
        protected_metrics_by_section=metrics,
        editable_paragraph_ids=[p.paragraph_id for p in paragraphs if p.editable],
        page_width_points=_points(section.page_width),
        page_height_points=_points(section.page_height),
        margins_points=margins,
    )


def _points(value: object) -> float:
    points = getattr(value, "pt", None)
    return float(points) if points is not None else 0.0


def _parse_paragraphs(document: DocxDocument) -> list[ResumeParagraph]:
    current_section = "header"
    current_entry = ""
    bullet_slots: dict[str, int] = {}
    parsed: list[ResumeParagraph] = []

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        upper = text.upper().strip()
        kind = ParagraphKind.protected_body

        if index == 0:
            kind = ParagraphKind.name
        elif index == 1:
            kind = ParagraphKind.contact
        elif not text:
            kind = ParagraphKind.spacer
        elif upper in SECTION_HEADINGS:
            current_section = upper.lower()
            current_entry = ""
            kind = ParagraphKind.section_heading
        elif current_section == "header":
            current_section = "summary"
            kind = ParagraphKind.summary
        elif current_section == "skills" and ":" in text:
            kind = ParagraphKind.skill_line
        elif current_section in {"experience", "projects"} and _entry_identity(text):
            current_entry = _entry_identity(text)
            kind = ParagraphKind.entry_heading
        elif current_entry:
            kind = ParagraphKind.bullet

        semantic_section = _semantic_section(current_section, current_entry, kind)
        bullet_slot: int | None = None
        if kind == ParagraphKind.bullet:
            bullet_slots[semantic_section] = bullet_slots.get(semantic_section, 0) + 1
            bullet_slot = bullet_slots[semantic_section]

        paragraph_id = _paragraph_id(kind, semantic_section, text, bullet_slot, index)
        run_models: list[TextRun] = []
        for item in paragraph.iter_inner_content():
            if isinstance(item, Hyperlink):
                for run in item.runs:
                    run_models.append(
                        TextRun(
                            text=run.text,
                            bold=run.bold,
                            italic=run.italic,
                            hyperlink_target=item.url or None,
                        )
                    )
            else:
                run_models.append(TextRun(text=item.text, bold=item.bold, italic=item.italic))

        editable = semantic_section in EDITABLE_SECTIONS and kind in {
            ParagraphKind.summary,
            ParagraphKind.skill_line,
            ParagraphKind.bullet,
        }
        parsed.append(
            ResumeParagraph(
                paragraph_id=paragraph_id,
                index=index,
                section=semantic_section,
                kind=kind,
                text=paragraph.text,
                runs=run_models,
                editable=editable,
                bullet_slot=bullet_slot,
                line_budget=_line_budget(kind, semantic_section, paragraph.text),
                character_budget=max(len(paragraph.text), 1),
            )
        )
    return parsed


def _entry_identity(text: str) -> str:
    upper = text.upper()
    for identity, entry_id in ENTRY_IDENTITIES.items():
        if upper.startswith(identity):
            return entry_id
    return ""


def _semantic_section(current: str, entry: str, kind: ParagraphKind) -> str:
    if kind in {ParagraphKind.name, ParagraphKind.contact, ParagraphKind.spacer}:
        return current
    if entry and current == "experience":
        return f"experience.{entry}"
    if entry and current == "projects":
        return f"projects.{entry}"
    return current


def _paragraph_id(
    kind: ParagraphKind,
    section: str,
    text: str,
    bullet_slot: int | None,
    index: int,
) -> str:
    if kind == ParagraphKind.summary:
        return "summary"
    if kind == ParagraphKind.skill_line:
        return f"skills.{slug(text.split(':', 1)[0])}"
    if kind == ParagraphKind.bullet:
        return f"{section}.bullet.{bullet_slot}"
    if kind == ParagraphKind.entry_heading:
        return f"{section}.heading"
    if kind == ParagraphKind.section_heading:
        return f"{section}.heading"
    if kind == ParagraphKind.name:
        return "identity.name"
    if kind == ParagraphKind.contact:
        return "identity.contact"
    return f"{section}.{kind.value}.{index}"


def _line_budget(kind: ParagraphKind, section: str, text: str) -> int:
    if kind == ParagraphKind.summary:
        return 3
    if kind == ParagraphKind.skill_line:
        return 1
    if kind == ParagraphKind.bullet:
        if section == "projects.loavenly" and len(text) < 125:
            return 1
        return 2
    return 1


def _build_sections(paragraphs: list[ResumeParagraph]) -> list[ResumeSection]:
    ordered: dict[str, list[ResumeParagraph]] = {}
    for paragraph in paragraphs:
        ordered.setdefault(paragraph.section, []).append(paragraph)
    result: list[ResumeSection] = []
    for section_id, members in ordered.items():
        heading = next(
            (p.text.strip() for p in members if p.kind == ParagraphKind.section_heading),
            section_id,
        )
        result.append(
            ResumeSection(
                section_id=section_id,
                heading=heading,
                paragraph_ids=[p.paragraph_id for p in members],
                editable_paragraph_ids=[p.paragraph_id for p in members if p.editable],
                bullet_count=sum(p.kind == ParagraphKind.bullet for p in members),
                line_budget=sum(p.line_budget for p in members),
            )
        )
    return result


def _protected_strings(paragraphs: list[ResumeParagraph]) -> list[str]:
    protected_kinds = {
        ParagraphKind.name,
        ParagraphKind.contact,
        ParagraphKind.entry_heading,
        ParagraphKind.section_heading,
        ParagraphKind.protected_body,
    }
    return [p.text for p in paragraphs if p.text.strip() and p.kind in protected_kinds]


def _hyperlinks(document: DocxDocument) -> list[str]:
    values: list[str] = []
    for paragraph in document.paragraphs:
        for item in paragraph.iter_inner_content():
            if isinstance(item, Hyperlink) and item.url:
                values.append(item.url)
    return values


def _metrics_by_section(paragraphs: list[ResumeParagraph]) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    for paragraph in paragraphs:
        if not paragraph.editable:
            continue
        matches: list[str] = []
        for pattern in METRIC_PATTERNS:
            matches.extend(match.group(0) for match in pattern.finditer(paragraph.text))
        if matches:
            result.setdefault(paragraph.section, []).extend(matches)
    return result
